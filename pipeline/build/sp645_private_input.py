from __future__ import annotations

import argparse
import hashlib
import json
from copy import deepcopy
from pathlib import Path

import openpyxl
import yaml

HISTORICAL_COLUMNS = [
    "receita_tributaria_pct_receita_corrente",
    "investimento_pct_receita_corrente",
    "despesa_territorial_pct_despesa",
    "dtp_pct_rcl",
    "dc_pct_rcl",
    "dcl_pct_rcl",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sheet_rows(path: Path, sheet_name: str) -> tuple[list[str], list[dict]]:
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        header = [str(value) if value is not None else "" for value in next(iterator)]
        rows = [dict(zip(header, values)) for values in iterator]
        return header, rows
    finally:
        book.close()


def prepare_sp645_private_input(
    typology_input: Path,
    multisource_2025: Path,
    base_mapping: Path,
    output_workbook: Path,
    output_mapping: Path,
    manifest_path: Path,
) -> dict:
    hist_header, hist_rows = _sheet_rows(typology_input, "Tipologia SP645")
    current_header, current_rows = _sheet_rows(multisource_2025, "Base multifuentes")

    if len(hist_rows) != 3225:
        raise ValueError(f"Expected 3225 typology rows, got {len(hist_rows)}")
    if len(current_rows) != 645:
        raise ValueError(f"Expected 645 current rows, got {len(current_rows)}")

    current_by_code = {str(row["cod_ibge"]).strip(): row for row in current_rows}
    if len(current_by_code) != 645:
        raise ValueError("Current workbook municipality codes are not unique.")

    years = sorted({int(row["ano"]) for row in hist_rows})
    if years != [2021, 2022, 2023, 2024, 2025]:
        raise ValueError(f"Unexpected history years: {years}")

    codes = {str(row["cod_ibge"]).strip() for row in hist_rows}
    if len(codes) != 645:
        raise ValueError(f"Expected 645 historical municipalities, got {len(codes)}")
    if codes != set(current_by_code):
        missing = sorted(codes - set(current_by_code))
        extra = sorted(set(current_by_code) - codes)
        raise ValueError(f"Universe mismatch. missing={missing[:10]} extra={extra[:10]}")

    for column in HISTORICAL_COLUMNS:
        if column not in hist_header:
            raise ValueError(f"Missing historical column: {column}")

    output_fields = list(current_header)
    for column in HISTORICAL_COLUMNS:
        if column not in output_fields:
            output_fields.append(column)

    output_rows: list[dict] = []
    for hist in hist_rows:
        code = str(hist["cod_ibge"]).strip()
        year = int(hist["ano"])
        row = {field: "" for field in output_fields}
        row["cod_ibge"] = code
        row["municipio"] = hist.get("municipio") or current_by_code[code].get("municipio")
        row["ano"] = year

        if year == 2025:
            for field in current_header:
                row[field] = current_by_code[code].get(field, "")

        for field in HISTORICAL_COLUMNS:
            value = hist.get(field)
            row[field] = "" if value is None else value

        output_rows.append(row)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "SP645 Site Input"
    sheet.append(output_fields)
    for row in output_rows:
        sheet.append([row.get(field, "") for field in output_fields])

    output_workbook.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_workbook)

    mapping = yaml.safe_load(base_mapping.read_text(encoding="utf-8"))
    mapping = deepcopy(mapping)
    mapping["source"].update(
        {
            "workbook_name": output_workbook.name,
            "sheet": "SP645 Site Input",
            "data_version": "sp645-private-site-2021-2025-v1",
            "expected_rows": 3225,
            "expected_municipalities": 645,
            "expected_years": {"start": 2021, "end": 2025},
        }
    )

    variables = mapping["variables"]
    for variable_id, start in (
        ("dtp_pct_rcl", 2023),
        ("dc_pct_rcl", 2023),
        ("dcl_pct_rcl", 2023),
    ):
        variables[variable_id]["column"] = variable_id
        variables[variable_id]["availability_start"] = start
        variables[variable_id]["availability_end"] = 2025
        variables[variable_id]["scale"] = 0.01

    variables["receita_tributaria_pct_receita_corrente"] = {
        "column": "receita_tributaria_pct_receita_corrente",
        "source_id": "SICONFI_DCA",
        "availability_start": 2021,
        "availability_end": 2025,
        "value_type": "numeric",
        "scale": 0.01,
    }
    variables["investimento_pct_receita_corrente"] = {
        "column": "investimento_pct_receita_corrente",
        "source_id": "SICONFI_DCA",
        "availability_start": 2021,
        "availability_end": 2025,
        "value_type": "numeric",
        "scale": 0.01,
    }
    variables["despesa_territorial_pct_despesa"] = {
        "column": "despesa_territorial_pct_despesa",
        "source_id": "SICONFI_DCA",
        "availability_start": 2021,
        "availability_end": 2025,
        "value_type": "numeric",
        "scale": 0.01,
    }

    output_mapping.parent.mkdir(parents=True, exist_ok=True)
    output_mapping.write_text(
        yaml.safe_dump(mapping, allow_unicode=True, sort_keys=False),
        encoding="utf-8",
    )

    result = {
        "rows": len(output_rows),
        "municipalities": len(codes),
        "years": years,
        "historical_columns": HISTORICAL_COLUMNS,
        "input_sha256": {
            "typology_input": sha256(typology_input),
            "multisource_2025": sha256(multisource_2025),
            "base_mapping": sha256(base_mapping),
        },
        "output_sha256": {
            "workbook": sha256(output_workbook),
            "mapping": sha256(output_mapping),
        },
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Prepare SP645 private-site analytical input.")
    parser.add_argument("--typology-input", type=Path, required=True)
    parser.add_argument("--multisource-2025", type=Path, required=True)
    parser.add_argument("--base-mapping", type=Path, required=True)
    parser.add_argument("--output-workbook", type=Path, required=True)
    parser.add_argument("--output-mapping", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()

    result = prepare_sp645_private_input(
        args.typology_input,
        args.multisource_2025,
        args.base_mapping,
        args.output_workbook,
        args.output_mapping,
        args.manifest,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
