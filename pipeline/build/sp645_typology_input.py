from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

import openpyxl

VARIABLES = [
    "receita_tributaria_pct_receita_corrente",
    "investimento_pct_receita_corrente",
    "despesa_territorial_pct_despesa",
    "dtp_pct_rcl",
    "dc_pct_rcl",
    "dcl_pct_rcl",
    "caixa_pos_rpnp_pct_rcl",
]

RGF_HISTORY_MAP = {
    "dtp_pct_rcl": "rgf_dtp_percentual_rcl",
    "dc_pct_rcl": "rgf02_dc_percentual_rcl",
    "dcl_pct_rcl": "rgf02_dcl_percentual_rcl",
}

YEARS = range(2021, 2026)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read_csv(path: Path) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _clean(value):
    if value is None:
        return None
    if isinstance(value, str) and not value.strip():
        return None
    return value


def build_typology_input(
    dca_csv: Path,
    rgf_history_csv: Path,
    gate_d_xlsx: Path,
    output_xlsx: Path,
    manifest_path: Path,
    *,
    expected_municipalities: int = 645,
) -> dict:
    dca_rows = _read_csv(dca_csv)
    history_rows = _read_csv(rgf_history_csv)

    dca = {}
    for row in dca_rows:
        key = (str(row["cod_ibge"]).strip(), int(row["ano"]))
        if key in dca:
            raise ValueError(f"Duplicate DCA key: {key}")
        dca[key] = row

    history = {}
    for row in history_rows:
        key = (str(row["cod_ibge"]).strip(), int(row["ano"]))
        if key in history:
            raise ValueError(f"Duplicate RGF history key: {key}")
        history[key] = row

    wb = openpyxl.load_workbook(gate_d_xlsx, read_only=True, data_only=True)
    try:
        ws = wb["Indicadores"]
        iterator = ws.iter_rows(values_only=True)
        headers = [str(value) for value in next(iterator)]
        idx = {name: i for i, name in enumerate(headers)}
        required = {
            "cod_ibge",
            "municipio",
            "ano",
            "dtp_pct_rcl",
            "dc_pct_rcl",
            "dcl_pct_rcl",
            "caixa_pos_rpnp_pct_rcl",
        }
        missing = sorted(required - set(idx))
        if missing:
            raise ValueError(f"Gate D columns missing: {missing}")
        gate_d = {}
        names = {}
        for values in iterator:
            code = str(values[idx["cod_ibge"]]).strip()
            year = int(values[idx["ano"]])
            if year != 2025:
                continue
            if code in gate_d:
                raise ValueError(f"Duplicate Gate D code: {code}")
            gate_d[code] = {
                field: _clean(values[idx[field]])
                for field in (
                    "dtp_pct_rcl",
                    "dc_pct_rcl",
                    "dcl_pct_rcl",
                    "caixa_pos_rpnp_pct_rcl",
                )
            }
            names[code] = str(values[idx["municipio"]]).strip()
    finally:
        wb.close()

    codes = sorted(names)
    if len(codes) != expected_municipalities:
        raise ValueError(
            f"Expected {expected_municipalities} municipalities in Gate D, got {len(codes)}"
        )

    expected_dca_keys = {(code, year) for code in codes for year in YEARS}
    if set(dca) != expected_dca_keys:
        missing = sorted(expected_dca_keys - set(dca))
        extra = sorted(set(dca) - expected_dca_keys)
        raise ValueError(
            f"DCA key mismatch: missing={missing[:10]} extra={extra[:10]}"
        )

    expected_history_keys = {(code, year) for code in codes for year in (2023, 2024)}
    if set(history) != expected_history_keys:
        missing = sorted(expected_history_keys - set(history))
        extra = sorted(set(history) - expected_history_keys)
        raise ValueError(
            f"RGF history key mismatch: missing={missing[:10]} extra={extra[:10]}"
        )

    rows = []
    for code in codes:
        for year in YEARS:
            row = {
                "cod_ibge": code,
                "municipio": names[code],
                "ano": year,
            }
            dca_row = dca[(code, year)]
            for field in (
                "receita_tributaria_pct_receita_corrente",
                "investimento_pct_receita_corrente",
                "despesa_territorial_pct_despesa",
            ):
                row[field] = _clean(dca_row.get(field))

            if year in (2023, 2024):
                history_row = history[(code, year)]
                for target, source in RGF_HISTORY_MAP.items():
                    row[target] = _clean(history_row.get(source))
            elif year == 2025:
                for field, value in gate_d[code].items():
                    row[field] = value

            rows.append(row)

    coverage = {}
    at_least_three = {}
    for field in VARIABLES:
        coverage[field] = {
            str(year): sum(
                row.get(field) is not None
                for row in rows
                if row["ano"] == year
            )
            for year in YEARS
        }
        by_code = {}
        for code in codes:
            by_code[code] = sum(
                row.get(field) is not None
                for row in rows
                if row["cod_ibge"] == code
            )
        at_least_three[field] = sum(n >= 3 for n in by_code.values())

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Tipologia SP645"
    headers = ["cod_ibge", "municipio", "ano", *VARIABLES]
    ws.append(headers)
    for row in rows:
        ws.append([row.get(field) for field in headers])
    out.save(output_xlsx)

    result = {
        "rows": len(rows),
        "municipalities": len(codes),
        "years": list(YEARS),
        "variables": VARIABLES,
        "coverage_by_year": coverage,
        "municipalities_with_at_least_3_years": at_least_three,
        "input_sha256": {
            "dca": sha256(dca_csv),
            "rgf_history": sha256(rgf_history_csv),
            "gate_d": sha256(gate_d_xlsx),
        },
        "output_sha256": sha256(output_xlsx),
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Build canonical SP645 annual input for typologies and pairs."
    )
    parser.add_argument("--dca", type=Path, required=True)
    parser.add_argument("--rgf-history", type=Path, required=True)
    parser.add_argument("--gate-d", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--expected-municipalities", type=int, default=645)
    args = parser.parse_args()

    print(
        json.dumps(
            build_typology_input(
                args.dca,
                args.rgf_history,
                args.gate_d,
                args.output,
                args.manifest,
                expected_municipalities=args.expected_municipalities,
            ),
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
