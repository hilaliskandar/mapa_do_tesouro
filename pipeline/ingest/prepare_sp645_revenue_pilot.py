from __future__ import annotations

import argparse
import hashlib
from pathlib import Path

import openpyxl

SOURCE_SHEET = "Base_Analitica"
OUTPUT_SHEET = "Base multifuentes"
YEARS = (2020, 2021, 2022, 2023)

SOURCE_COLUMNS = {
    "cod_ibge": "cod_ibge",
    "municipio": "municipio",
    "ano": "ano_finbra",
    "sp645_receitas_correntes": "receitas_correntes",
    "sp645_receita_tributaria_total": "receita_tributaria_total",
    "sp645_iptu": "iptu",
    "sp645_itbi": "itbi",
    "sp645_iss": "iss",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_ibge(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    code = str(value).strip()
    if len(code) != 7 or not code.isdigit():
        raise ValueError(f"Codigo IBGE invalido: {value!r}")
    return code


def prepare_sp645_revenue_pilot(
    source: Path,
    output: Path,
    *,
    expected_municipalities: int = 645,
    expected_years: tuple[int, ...] = YEARS,
) -> dict:
    source_hash = sha256(source)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if SOURCE_SHEET not in workbook.sheetnames:
            raise ValueError(f"Aba obrigatoria ausente: {SOURCE_SHEET}")

        sheet = workbook[SOURCE_SHEET]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value) if value is not None else "" for value in next(rows)]
        indexes = {name: index for index, name in enumerate(headers)}
        missing = sorted(set(SOURCE_COLUMNS.values()) - set(indexes))
        if missing:
            raise ValueError(f"Colunas obrigatorias ausentes: {missing}")

        records = []
        seen = set()
        municipalities = {}
        for row_number, row in enumerate(rows, start=2):
            code = normalize_ibge(row[indexes[SOURCE_COLUMNS["cod_ibge"]]])
            name = str(row[indexes[SOURCE_COLUMNS["municipio"]]]).strip()
            year = int(row[indexes[SOURCE_COLUMNS["ano"]]])

            if year not in expected_years:
                raise ValueError(f"Exercicio inesperado na linha {row_number}: {year}")
            key = (code, year)
            if key in seen:
                raise ValueError(f"Par municipio-ano duplicado: {key}")
            seen.add(key)

            previous_name = municipalities.setdefault(code, name)
            if previous_name != name:
                raise ValueError(f"Nome municipal inconsistente para {code}")

            record = {
                output_column: row[indexes[source_column]]
                for output_column, source_column in SOURCE_COLUMNS.items()
            }
            record["cod_ibge"] = code
            record["municipio"] = name
            record["ano"] = year
            records.append(record)
    finally:
        workbook.close()

    expected_rows = expected_municipalities * len(expected_years)
    if len(municipalities) != expected_municipalities:
        raise ValueError(
            f"Quantidade de municipios inesperada: {len(municipalities)} != "
            f"{expected_municipalities}"
        )
    if len(records) != expected_rows:
        raise ValueError(
            f"Quantidade de linhas inesperada: {len(records)} != {expected_rows}"
        )
    observed_years = sorted({record["ano"] for record in records})
    if observed_years != list(expected_years):
        raise ValueError(
            f"Exercicios inesperados: {observed_years} != {list(expected_years)}"
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    out = openpyxl.Workbook()
    base = out.active
    base.title = OUTPUT_SHEET
    columns = list(SOURCE_COLUMNS)
    base.append(columns)
    for record in sorted(records, key=lambda item: (item["cod_ibge"], item["ano"])):
        base.append([record[column] for column in columns])

    manifest = out.create_sheet("Manifesto")
    manifest.append(["campo", "valor"])
    manifest.append(["source_sha256", source_hash])
    manifest.append(["source_sheet", SOURCE_SHEET])
    manifest.append(["rows", len(records)])
    manifest.append(["municipalities", len(municipalities)])
    manifest.append(["years", ",".join(str(year) for year in observed_years)])
    manifest.append([
        "scope",
        "Piloto estadual de receitas; nao equivale ao baseline multifuentes 2013-2025.",
    ])
    manifest.append([
        "excluded",
        "Despesas empenhadas reconstruidas, populacao_2024 e indicadores derivados.",
    ])
    out.save(output)
    out.close()

    return {
        "source_sha256": source_hash,
        "rows": len(records),
        "municipalities": len(municipalities),
        "years": observed_years,
        "variables": len(SOURCE_COLUMNS) - 3,
        "output": str(output),
    }


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Prepara o piloto estadual de receitas para 645 municipios."
    )
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    result = prepare_sp645_revenue_pilot(args.source, args.output)
    for key, value in result.items():
        print(f"{key}={value}")


if __name__ == "__main__":
    main()
