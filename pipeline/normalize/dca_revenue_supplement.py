from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from pipeline.normalize.dca import DEFAULT_MAPPING, load_mapping, normalize_bundle


EXPECTED_COLUMNS = (
    "Ano",
    "Instituição",
    "Cod.IBGE",
    "UF",
    "População",
    "Coluna",
    "Conta",
    "Identificador da Conta",
    "Valor",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean_header(value: object) -> str:
    return str(value or "").strip()


def _parse_code(value: object) -> str:
    if value is None or str(value).strip() == "":
        raise ValueError("Missing municipality code")
    if isinstance(value, (int, float)):
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _parse_population(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    return int(float(str(value).strip().replace(",", ".")))


def _parse_value(value: object) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text)


def normalize_revenue_supplement(
    workbook_path: Path,
    year: int,
    codes: list[str],
    output_csv: Path,
    report_path: Path,
    *,
    mapping_path: Path = DEFAULT_MAPPING,
) -> dict:
    requested = []
    seen = set()
    for raw in codes:
        code = str(raw).strip()
        if not code:
            continue
        if len(code) != 7 or not code.isdigit():
            raise ValueError(f"Invalid 7-digit IBGE code: {code!r}")
        if code not in seen:
            seen.add(code)
            requested.append(code)
    if not requested:
        raise ValueError("At least one municipality code is required.")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet_name = f"{int(year)}-finbra"
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(
            f"Worksheet {sheet_name!r} not found in {workbook_path}. "
            f"Available: {workbook.sheetnames}"
        )
    worksheet = workbook[sheet_name]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        header = [_clean_header(value) for value in next(iterator)]
    except StopIteration as exc:
        workbook.close()
        raise ValueError(f"Empty worksheet: {sheet_name}") from exc

    if tuple(header) != EXPECTED_COLUMNS:
        workbook.close()
        raise ValueError(
            f"Unexpected header in {sheet_name}: {header!r}; "
            f"expected {list(EXPECTED_COLUMNS)!r}"
        )

    index = {name: position for position, name in enumerate(header)}
    requested_set = set(requested)
    by_code: dict[str, list[dict]] = defaultdict(list)
    names: dict[str, str] = {}
    row_counts: dict[str, int] = defaultdict(int)

    for row in iterator:
        if not row or len(row) < len(header):
            continue
        code = _parse_code(row[index["Cod.IBGE"]])
        if code not in requested_set:
            continue

        row_year = int(row[index["Ano"]])
        if row_year != int(year):
            workbook.close()
            raise ValueError(
                f"Unexpected exercise in {sheet_name}: {row_year} != {year}"
            )

        identifier = str(row[index["Identificador da Conta"]] or "").strip()
        by_code[code].append(
            {
                "entity_id": code,
                "year": int(year),
                "annex": "DCA-Anexo I-C",
                "populacao": _parse_population(row[index["População"]]),
                "coluna": str(row[index["Coluna"]] or "").strip(),
                "conta": str(row[index["Conta"]] or "").strip(),
                "cod_conta": identifier.removeprefix("siconfi-cor_"),
                "valor": _parse_value(row[index["Valor"]]),
                "rotulo": "Padrão",
            }
        )
        row_counts[code] += 1
        names.setdefault(code, str(row[index["Instituição"]] or "").strip())

    workbook.close()

    missing_codes = sorted(requested_set - set(by_code))
    if missing_codes:
        raise ValueError(
            f"Requested municipality codes not found in {sheet_name}: {missing_codes}"
        )

    mapping = load_mapping(mapping_path)
    variable_ids = list(mapping["variables"])
    revenue_mapping = {
        "version": mapping.get("version"),
        "variables": {
            variable_id: rule
            for variable_id, rule in mapping["variables"].items()
            if rule.get("annex") == "DCA-Anexo I-C"
            or rule.get("special") == "population_consensus"
        },
    }

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    issue_count = 0
    observed_counts = {variable_id: 0 for variable_id in variable_ids}

    with output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["cod_ibge", "ano", *variable_ids, "qa_issue_count"],
        )
        writer.writeheader()

        for code in requested:
            payload = {
                "DCA-Anexo I-C": {
                    "entity_id": code,
                    "year": int(year),
                    "annex": "DCA-Anexo I-C",
                    "items": by_code[code],
                }
            }
            normalized = normalize_bundle(payload, mapping=revenue_mapping)
            issue_count += len(normalized["issues"])

            row = {
                "cod_ibge": code,
                "ano": int(year),
                "qa_issue_count": len(normalized["issues"]),
            }
            for variable_id in variable_ids:
                entry = normalized["variables"].get(variable_id)
                if entry and entry["status"] == "observado":
                    row[variable_id] = entry["value"]
                    observed_counts[variable_id] += 1
                else:
                    row[variable_id] = ""
            writer.writerow(row)

    report = {
        "source": "FINBRA_DCA_CONSOLIDATED_REVENUE_WORKBOOK",
        "workbook": str(workbook_path),
        "workbook_sha256": sha256(workbook_path),
        "worksheet": sheet_name,
        "year": int(year),
        "codes": requested,
        "municipalities": [
            {
                "code": code,
                "name": names.get(code, ""),
                "source_rows": row_counts[code],
            }
            for code in requested
        ],
        "rows": len(requested),
        "issues": issue_count,
        "observed_counts": {
            variable_id: count
            for variable_id, count in observed_counts.items()
            if count
        },
        "output_csv": str(output_csv),
        "output_sha256": sha256(output_csv),
        "policy": "local_fill_only_supplement",
    }

    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Normalize selected municipality revenues from the consolidated "
            "FINBRA 2022-2025 workbook as a local fill-only supplement."
        )
    )
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--year", type=int, required=True)
    parser.add_argument("--codes", nargs="+", required=True)
    parser.add_argument("--output-csv", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING)
    args = parser.parse_args()

    result = normalize_revenue_supplement(
        args.workbook,
        args.year,
        args.codes,
        args.output_csv,
        args.report,
        mapping_path=args.mapping,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
