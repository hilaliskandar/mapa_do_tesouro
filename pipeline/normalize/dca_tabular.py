from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from pipeline.acquire.siconfi_dca import load_codes
from pipeline.normalize.dca import DEFAULT_MAPPING, load_mapping, normalize_bundle

REQUIRED_COLUMNS = (
    "Instituição",
    "Cod.IBGE",
    "UF",
    "População",
    "Coluna",
    "Conta",
    "Identificador da Conta",
    "Valor",
)

ANNEXES = (
    "DCA-Anexo I-C",
    "DCA-Anexo I-D",
    "DCA-Anexo I-E",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _parse_year(value: object) -> int:
    match = re.search(r"(20\d{2})", str(value or ""))
    if not match:
        raise ValueError(f"Unable to parse exercise from metadata: {value!r}")
    return int(match.group(1))


def _parse_annex(value: object) -> str:
    text = str(value or "")
    for suffix in ("I-C", "I-D", "I-E"):
        if f"Anexo {suffix}" in text:
            return f"DCA-Anexo {suffix}"
    raise ValueError(f"Unable to parse DCA annex from metadata: {value!r}")


def _parse_code(value: object) -> str:
    if value is None or str(value).strip() == "":
        raise ValueError("Missing municipality code")
    if isinstance(value, (int, float)):
        return str(int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d+(?:\.0+)?", text):
        return str(int(float(text)))
    return text


def _parse_population(value: object) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return int(value)
    return int(float(str(value).strip().replace(",", ".")))


def _parse_value(value: object) -> float | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    return float(text)


def _csv_rows(path: Path) -> tuple[list[tuple], list[str]]:
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                reader = csv.reader(handle, delimiter=";")
                metadata = [tuple(next(reader)) for _ in range(3)]
                header = list(next(reader))
                rows = [tuple(row) for row in reader if row]
            return metadata + rows, header
        except (UnicodeDecodeError, StopIteration) as exc:
            last_error = exc
    raise ValueError(f"Unable to read CSV {path}: {last_error}")


def _xlsx_rows(path: Path) -> tuple[list[tuple], list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    iterator = worksheet.iter_rows(values_only=True)
    try:
        metadata = [tuple(next(iterator)) for _ in range(3)]
        header = [str(value) if value is not None else "" for value in next(iterator)]
    except StopIteration as exc:
        raise ValueError(f"Incomplete XLSX structure: {path}") from exc
    rows = [tuple(row) for row in iterator if any(value is not None for value in row)]
    workbook.close()
    return metadata + rows, header


def read_tabular_annex(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        combined, header = _csv_rows(path)
    elif suffix in {".xlsx", ".xlsm"}:
        combined, header = _xlsx_rows(path)
    else:
        raise ValueError(f"Unsupported tabular DCA format: {path.suffix}")

    metadata = combined[:3]
    rows = combined[3:]
    if len(metadata) != 3:
        raise ValueError(f"Missing metadata rows in {path}")
    year = _parse_year(metadata[0][0] if metadata[0] else None)
    annex = _parse_annex(metadata[2][0] if metadata[2] else None)

    if tuple(header) != REQUIRED_COLUMNS:
        raise ValueError(
            f"Unexpected header in {path}: {header!r}; "
            f"expected {list(REQUIRED_COLUMNS)!r}"
        )

    index = {name: position for position, name in enumerate(header)}
    by_code: dict[str, list[dict]] = defaultdict(list)
    names: dict[str, str] = {}

    for row in rows:
        if len(row) < len(header):
            continue
        code = _parse_code(row[index["Cod.IBGE"]])
        identifier = str(row[index["Identificador da Conta"]] or "").strip()
        cod_conta = identifier.removeprefix("siconfi-cor_")
        population = _parse_population(row[index["População"]])
        account = str(row[index["Conta"]] or "")
        institution = str(row[index["Instituição"]] or "")
        names.setdefault(code, institution)
        by_code[code].append(
            {
                "entity_id": code,
                "year": year,
                "annex": annex,
                "populacao": population,
                "coluna": str(row[index["Coluna"]] or ""),
                "conta": account,
                "cod_conta": cod_conta,
                "valor": _parse_value(row[index["Valor"]]),
                "rotulo": (
                    "Total Geral da Despesa por Função"
                    if annex == "DCA-Anexo I-E"
                    else "Padrão"
                ),
            }
        )

    return {
        "path": str(path),
        "year": year,
        "annex": annex,
        "codes": sorted(by_code),
        "by_code": dict(by_code),
        "names": names,
        "sha256": sha256(path),
        "size_bytes": path.stat().st_size,
        "format": suffix.lstrip("."),
    }


def _write_normalized(
    *,
    sources: list[dict],
    expected_codes: list[str],
    output_csv: Path,
    mapping_path: Path,
) -> dict:
    mapping = load_mapping(mapping_path)
    variable_ids = list(mapping["variables"])
    source_by_annex = {source["annex"]: source for source in sources}
    year = sources[0]["year"]
    issue_count = 0
    observed_counts = {variable_id: 0 for variable_id in variable_ids}

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    with output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["cod_ibge", "ano", *variable_ids, "qa_issue_count"],
        )
        writer.writeheader()
        for code in expected_codes:
            payloads = {}
            for annex in ANNEXES:
                source = source_by_annex[annex]
                payloads[annex] = {
                    "entity_id": code,
                    "year": year,
                    "annex": annex,
                    "items": source["by_code"].get(code, []),
                }
            normalized = normalize_bundle(payloads, mapping=mapping)
            issue_count += len(normalized["issues"])
            row = {
                "cod_ibge": code,
                "ano": year,
                "qa_issue_count": len(normalized["issues"]),
            }
            for variable_id in variable_ids:
                entry = normalized["variables"][variable_id]
                if entry["status"] == "observado":
                    row[variable_id] = entry["value"]
                    observed_counts[variable_id] += 1
                else:
                    row[variable_id] = ""
            writer.writerow(row)

    return {
        "rows": len(expected_codes),
        "variables": len(variable_ids),
        "issues": issue_count,
        "observed_counts": observed_counts,
        "output": str(output_csv),
    }


def normalize_tabular_files(
    paths: Iterable[Path],
    output_csv: Path,
    *,
    universe_geojson: Path | None = None,
    mapping_path: Path = DEFAULT_MAPPING,
    report_path: Path | None = None,
) -> dict:
    sources = [read_tabular_annex(Path(path)) for path in paths]
    if len(sources) != 3:
        raise ValueError("Exactly three DCA tabular files are required: I-C, I-D and I-E")

    years = {source["year"] for source in sources}
    if len(years) != 1:
        raise ValueError(f"Mixed exercises are not allowed: {sorted(years)}")
    annexes = [source["annex"] for source in sources]
    if set(annexes) != set(ANNEXES) or len(set(annexes)) != 3:
        raise ValueError(f"Expected one file for each annex {ANNEXES}; got {annexes}")

    year = next(iter(years))
    source_by_annex = {source["annex"]: source for source in sources}
    present_union = set().union(*(set(source["codes"]) for source in sources))
    present_intersection = set.intersection(
        *(set(source["codes"]) for source in sources)
    )

    if universe_geojson is not None:
        expected_codes = [code for code, _ in load_codes(universe_geojson)]
    else:
        expected_codes = sorted(present_union)
    expected_set = set(expected_codes)

    normalization = _write_normalized(
        sources=sources,
        expected_codes=expected_codes,
        output_csv=output_csv,
        mapping_path=mapping_path,
    )

    report = {
        "source": "FINBRA_DCA_TABULAR_EXPORT",
        "year": year,
        "expected_municipalities": len(expected_codes),
        "municipalities_present_any_annex": len(present_union),
        "municipalities_present_all_annexes": len(present_intersection),
        "unexpected_codes": sorted(present_union - expected_set),
        "annexes": {
            annex: {
                "path": source_by_annex[annex]["path"],
                "format": source_by_annex[annex]["format"],
                "size_bytes": source_by_annex[annex]["size_bytes"],
                "sha256": source_by_annex[annex]["sha256"],
                "municipalities_present": len(source_by_annex[annex]["codes"]),
                "missing_codes": sorted(
                    expected_set - set(source_by_annex[annex]["codes"])
                ),
            }
            for annex in ANNEXES
        },
        "normalization": normalization,
        "coverage": {
            variable_id: {
                "observed": observed,
                "missing": len(expected_codes) - observed,
                "coverage_pct": (
                    round(observed / len(expected_codes) * 100, 2)
                    if expected_codes
                    else 0.0
                ),
            }
            for variable_id, observed in normalization["observed_counts"].items()
        },
    }

    if report_path is not None:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Normalize official tabular FINBRA/DCA exports (CSV or XLSX) "
            "without re-querying the SICONFI API."
        )
    )
    parser.add_argument("--i-c", type=Path, required=True)
    parser.add_argument("--i-d", type=Path, required=True)
    parser.add_argument("--i-e", type=Path, required=True)
    parser.add_argument("--output-csv", type=Path, required=True)
    parser.add_argument("--universe-geojson", type=Path)
    parser.add_argument("--mapping", type=Path, default=DEFAULT_MAPPING)
    parser.add_argument("--report", type=Path)
    args = parser.parse_args()

    result = normalize_tabular_files(
        [args.i_c, args.i_d, args.i_e],
        args.output_csv,
        universe_geojson=args.universe_geojson,
        mapping_path=args.mapping,
        report_path=args.report,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
