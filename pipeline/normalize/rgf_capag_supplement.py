from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


FIELD_HEADERS = {
    "rgf_caixa_bruta_nao_vinculada":
        "2025DisponibilidadeDeCaixaBrutaDISPONIBILIDADE DE CAIXA BRUTA (a)"
        "TOTAL DOS RECURSOS NÃO VINCULADOS (I)",
    "rgf_demais_obrigacoes_nao_vinculadas":
        "2025DemaisObrigacoesFinanceirasDemais Obrigações Financeiras (e)"
        "TOTAL DOS RECURSOS NÃO VINCULADOS (I)",
    "rgf_rp_nao_liquidados_anteriores_nao_vinculados":
        "2025RestosAPagarEmpenhadosENaoLiquidadosDeExerciciosAnteriores"
        "Restos a Pagar Empenhados e Não Liquidados de Exercícios Anteriores (d)"
        "TOTAL DOS RECURSOS NÃO VINCULADOS (I)",
    "rgf_rp_liquidados_anteriores_nao_vinculados":
        "2025RestosAPagarLiquidadosENaoPagosDeExerciciosAnteriores"
        "De Exercícios Anteriores (b)TOTAL DOS RECURSOS NÃO VINCULADOS (I)",
    "rgf_rp_liquidados_exercicio_nao_vinculados":
        "2025RestosAPagarLiquidadosENaoPagosDoExercicio"
        "Do Exercício (c)TOTAL DOS RECURSOS NÃO VINCULADOS (I)",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_supplement(
    rgf_csv: Path,
    capag_xlsx: Path,
    output_csv: Path,
    manifest_path: Path,
) -> dict:
    with rgf_csv.open(encoding="utf-8", newline="") as handle:
        base_rows = list(csv.DictReader(handle))
    if not base_rows:
        raise ValueError("RGF CSV is empty.")

    base = {row["cod_ibge"]: row for row in base_rows}
    if len(base) != len(base_rows):
        raise ValueError("Duplicate cod_ibge in RGF CSV.")

    workbook = load_workbook(capag_xlsx, read_only=True, data_only=True)
    sheet = workbook["Datalake"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    header_index = {str(value): index for index, value in enumerate(header) if value is not None}

    missing_headers = [value for value in FIELD_HEADERS.values() if value not in header_index]
    if missing_headers:
        raise ValueError(f"Missing CAPAG Datalake headers: {missing_headers}")

    columns = {field: header_index[value] for field, value in FIELD_HEADERS.items()}
    capag: dict[str, dict[str, float | None]] = {}
    for row in sheet.iter_rows(min_row=3, values_only=True):
        try:
            code = str(int(row[0]))
        except (TypeError, ValueError):
            continue
        if code not in base:
            continue
        values = {}
        for field, column_index in columns.items():
            raw = row[column_index]
            values[field] = None if raw in (None, "") else float(raw)
        capag[code] = values

    output_csv.parent.mkdir(parents=True, exist_ok=True)
    output_fields = [
        "cod_ibge",
        "ano",
        *FIELD_HEADERS.keys(),
        "supplement_cell_count",
    ]

    fills: list[tuple[str, str, float]] = []
    conflicts: list[dict] = []
    with output_csv.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=output_fields)
        writer.writeheader()

        for code in sorted(base):
            base_row = base[code]
            local = capag.get(code, {})
            output_row = {"cod_ibge": code, "ano": base_row["ano"]}
            cell_count = 0

            for field in FIELD_HEADERS:
                base_value = base_row.get(field, "")
                local_value = local.get(field)

                if base_value not in ("", None):
                    if (
                        local_value is not None
                        and abs(float(base_value) - float(local_value)) > 0.01
                    ):
                        conflicts.append(
                            {
                                "cod_ibge": code,
                                "field": field,
                                "api_value": float(base_value),
                                "capag_value": float(local_value),
                            }
                        )
                    output_row[field] = ""
                elif local_value is not None:
                    output_row[field] = local_value
                    cell_count += 1
                    fills.append((code, field, local_value))
                else:
                    output_row[field] = ""

            output_row["supplement_cell_count"] = cell_count
            writer.writerow(output_row)

    manifest = {
        "rows": len(base),
        "filled_cells": len(fills),
        "conflicts": len(conflicts),
        "fills_by_field": {
            field: sum(1 for _, filled_field, _ in fills if filled_field == field)
            for field in FIELD_HEADERS
        },
        "source_rgf_sha256": sha256(rgf_csv),
        "source_capag_sha256": sha256(capag_xlsx),
        "output_sha256": sha256(output_csv),
        "conflict_details": conflicts,
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description="Build fill-only RGF supplement from CAPAG.")
    parser.add_argument("--rgf-csv", type=Path, required=True)
    parser.add_argument("--capag-xlsx", type=Path, required=True)
    parser.add_argument("--output-csv", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()

    result = build_supplement(
        args.rgf_csv,
        args.capag_xlsx,
        args.output_csv,
        args.manifest,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result["conflicts"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
