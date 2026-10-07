from __future__ import annotations

import argparse
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

AGGREGATIONS = {
    "tributos_imobiliarios": ["dca_iptu_principal", "dca_itbi_principal"],
    "tributos_selecionados": [
        "dca_iptu_principal", "dca_itbi_principal", "dca_iss_principal"
    ],
    "transferencias_selecionadas": [
        "dca_fpm_cota_mensal", "dca_icms_cota_parte", "dca_ipva_cota_parte"
    ],
    "despesa_territorial": [
        "dca_func_urbanismo_liquidada",
        "dca_func_habitacao_liquidada",
        "dca_func_saneamento_liquidada",
        "dca_func_gestao_ambiental_liquidada",
        "dca_func_transporte_liquidada",
    ],
    "servico_divida_liquidado": [
        "dca_juros_encargos_liquidada",
        "dca_amortizacao_divida_liquidada",
    ],
    "receitas_capital_selecionadas": [
        "dca_operacoes_credito",
        "dca_alienacao_bens",
        "dca_transferencias_capital",
    ],
    "despesas_capital_selecionadas": [
        "dca_investimentos_liquidada",
        "dca_inversoes_financeiras_liquidada",
        "dca_amortizacao_divida_liquidada",
    ],
    "gasto_social_selecionado": [
        "dca_func_saude_liquidada",
        "dca_func_educacao_liquidada",
    ],
}

MATERIALIZED_V7 = [
    "servico_divida_liquidado",
    "receitas_capital_selecionadas",
    "despesas_capital_selecionadas",
    "gasto_social_selecionado",
    "saldo_corrente_simplificado",
    "diferenca_capital_selecionada",
]


def load_rows(path: Path, sheet: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return header, rows


def equal(a, b):
    if a is None or b is None:
        return a is None and b is None
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return math.isclose(float(a), float(b), rel_tol=1e-12, abs_tol=1e-6)
    return a == b


def strict_sum(row, components):
    values = [row.get(x) for x in components]
    if not all(v is not None for v in values):
        return None
    return sum(float(v) for v in values)


def calculate(row):
    out = {key: strict_sum(row, parts) for key, parts in AGGREGATIONS.items()}

    a = row.get("dca_receita_corrente_bruta")
    b = row.get("dca_despesa_corrente_liquidada")
    out["saldo_corrente_simplificado"] = (
        None if a is None or b is None else float(a) - float(b)
    )

    a = out["receitas_capital_selecionadas"]
    b = out["despesas_capital_selecionadas"]
    out["diferenca_capital_selecionada"] = (
        None if a is None or b is None else a - b
    )
    return out


def compare(base: Path, panel_v7: Path, block3: Path) -> dict:
    base_header, base_rows = load_rows(base, "Base multifuentes")
    panel_header, panel_rows = load_rows(panel_v7, "Dados anuais")

    source = {
        (str(r["cod_ibge"]), int(r["ano"])): r
        for r in base_rows
    }
    panel = {
        (str(r["cod_ibge"]), int(r["ano"])): r
        for r in panel_rows
    }

    shared_fields = [
        field
        for field in base_header
        if field in panel_header and field != "municipio"
    ]
    source_mismatches = []
    for key, row in source.items():
        for field in shared_fields:
            if not equal(row.get(field), panel[key].get(field)):
                source_mismatches.append(
                    {
                        "codigo_ibge": key[0],
                        "ano": key[1],
                        "campo": field,
                        "base": row.get(field),
                        "v7": panel[key].get(field),
                    }
                )

    calculated = {key: calculate(row) for key, row in source.items()}
    aggregate_parity = {}
    for variable_id in MATERIALIZED_V7:
        mismatch = []
        observed_calc = observed_v7 = 0
        for key in source:
            a = calculated[key].get(variable_id)
            b = panel[key].get(variable_id)
            observed_calc += a is not None
            observed_v7 += b is not None
            if not equal(a, b):
                mismatch.append(
                    {
                        "codigo_ibge": key[0],
                        "ano": key[1],
                        "calculado": a,
                        "v7": b,
                    }
                )
        aggregate_parity[variable_id] = {
            "observado_calculado": observed_calc,
            "observado_v7": observed_v7,
            "divergencias": len(mismatch),
            "exemplos": mismatch[:10],
        }

    functions = AGGREGATIONS["despesa_territorial"]
    component_counts = Counter()
    missing_patterns = Counter()
    examples = []
    for key, row in source.items():
        values = [row.get(x) for x in functions]
        observed = sum(v is not None for v in values)
        component_counts[observed] += 1
        if 0 < observed < len(functions):
            missing = tuple(
                field for field, value in zip(functions, values)
                if value is None
            )
            missing_patterns[missing] += 1
            if len(examples) < 30:
                examples.append(
                    {
                        "codigo_ibge": key[0],
                        "municipio": row["municipio"],
                        "ano": key[1],
                        "componentes_observados": observed,
                        "ausentes": list(missing),
                    }
                )

    wb = load_workbook(block3, read_only=True, data_only=False)
    ws = wb["Indicadores anuais"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    first = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    formulas = {headers[i]: first[i] for i in range(29, len(headers))}

    return {
        "universe": {
            "rows": len(source),
            "municipalities": len({key[0] for key in source}),
            "years": sorted({key[1] for key in source}),
        },
        "source_parity": {
            "shared_fields": len(shared_fields),
            "cells_compared": len(source) * len(shared_fields),
            "mismatches": len(source_mismatches),
            "examples": source_mismatches[:20],
        },
        "materialized_v7_aggregate_parity": aggregate_parity,
        "territorial_rule_conflict": {
            "strict_complete_rows": component_counts[len(functions)],
            "legacy_nonempty_rows": sum(
                count for observed, count in component_counts.items()
                if observed > 0
            ),
            "rows_affected_by_partial_sum": sum(
                count for observed, count in component_counts.items()
                if 0 < observed < len(functions)
            ),
            "observed_component_distribution": dict(
                sorted(component_counts.items())
            ),
            "top_missing_patterns": [
                {"missing": list(pattern), "rows": count}
                for pattern, count in missing_patterns.most_common(15)
            ],
            "examples": examples,
            "block3_formula": formulas.get("despesa_territorial"),
            "canonical_rule": (
                "all five territorial components required; otherwise NA"
            ),
        },
        "block3_formula_baseline": formulas,
    }


def main():
    parser = argparse.ArgumentParser(
        description="Compara base multifuentes, painel v7 e Bloco 3."
    )
    parser.add_argument("base", type=Path)
    parser.add_argument("panel_v7", type=Path)
    parser.add_argument("block3", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    report = compare(args.base, args.panel_v7, args.block3)
    args.output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
