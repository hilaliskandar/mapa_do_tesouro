from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

import openpyxl

from pipeline.normalize.rgf02 import COLUMN as RGF02_COLUMN

A01_MAP = {
    "rgf_despesa_total_pessoal": "dtp_legal",
    "rgf_rcl_denominador_legal": "rcl_legal_ajustada",
    "rgf_dtp_percentual_rcl": "dtp_pct_rcl",
}

A02_CODES = {
    "rgf02_divida_consolidada": "DividaConsolidada",
    "rgf02_divida_consolidada_liquida": "DividaConsolidadaLiquida",
    "rgf02_dc_percentual_rcl": "PercentualDaDCSobreARCL",
    "rgf02_dcl_percentual_rcl": "PercentualDaDCLSobreARCL",
}

PERCENT_FIELDS = {
    "rgf_dtp_percentual_rcl",
    "rgf02_dc_percentual_rcl",
    "rgf02_dcl_percentual_rcl",
}


def _norm_code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _num(value):
    if value is None or value == "":
        return None
    return float(value)


def read_candidate(path: Path) -> dict[tuple[str, int], dict]:
    rows = {}
    with path.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            key = (_norm_code(row["cod_ibge"]), int(row["ano"]))
            if key in rows:
                raise ValueError(f"Duplicate candidate key: {key}")
            rows[key] = row
    return rows


def read_a01_reference(path: Path) -> dict[tuple[str, int], dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {}
    try:
        for year in (2023, 2024):
            ws = wb[f"RGF_{year}"]
            iterator = ws.iter_rows(values_only=True)
            headers = [str(value) for value in next(iterator)]
            idx = {name: i for i, name in enumerate(headers)}
            for values in iterator:
                if values[idx["anexo"]] != "RGF-Anexo 01":
                    continue
                code = _norm_code(values[idx["cod_ibge"]])
                row = result.setdefault((code, year), {})
                for target, source in A01_MAP.items():
                    row[target] = _num(values[idx[source]])
    finally:
        wb.close()
    return result


def read_a02_reference(path: Path) -> dict[tuple[str, int], dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {}
    reverse = {code: field for field, code in A02_CODES.items()}
    try:
        ws = wb["RGF02_LONGA"]
        iterator = ws.iter_rows(values_only=True)
        headers = [str(value) for value in next(iterator)]
        idx = {name: i for i, name in enumerate(headers)}
        for values in iterator:
            year = int(values[idx["ano"]])
            if year not in (2023, 2024):
                continue
            if values[idx["coluna"]] != RGF02_COLUMN:
                continue
            field = reverse.get(values[idx["cod_conta"]])
            if field is None:
                continue
            code = _norm_code(values[idx["cod_ibge"]])
            row = result.setdefault((code, year), {})
            value = _num(values[idx["valor"]])
            if field in row and row[field] != value:
                raise ValueError(
                    f"Reference A02 conflict: {(code, year, field)} "
                    f"{row[field]} != {value}"
                )
            row[field] = value
    finally:
        wb.close()
    return result


def compare_history(
    candidate_csv: Path,
    reference_a01_xlsx: Path,
    reference_a02_xlsx: Path,
    output_json: Path,
    *,
    money_tolerance: float = 0.01,
    percentage_tolerance: float = 0.01,
) -> dict:
    candidate = read_candidate(candidate_csv)
    a01 = read_a01_reference(reference_a01_xlsx)
    a02 = read_a02_reference(reference_a02_xlsx)

    reference = {}
    for key in set(a01) | set(a02):
        reference[key] = {**a01.get(key, {}), **a02.get(key, {})}

    codes = sorted({code for code, _year in reference})
    years = sorted({year for _code, year in reference})
    mismatches = []
    comparisons = 0
    exact_or_tolerant = 0
    both_absent = 0

    for key in sorted(reference, key=lambda item: (item[1], item[0])):
        if key not in candidate:
            mismatches.append(
                {
                    "cod_ibge": key[0],
                    "ano": key[1],
                    "field": "__row__",
                    "type": "candidate_row_missing",
                }
            )
            continue

        for field in [*A01_MAP, *A02_CODES]:
            ref = reference[key].get(field)
            cand = _num(candidate[key].get(field))
            comparisons += 1

            if ref is None and cand is None:
                both_absent += 1
                continue
            if ref is None and cand is not None:
                mismatches.append(
                    {
                        "cod_ibge": key[0],
                        "ano": key[1],
                        "field": field,
                        "type": "candidate_extra_value",
                        "reference": None,
                        "candidate": cand,
                    }
                )
                continue
            if ref is not None and cand is None:
                mismatches.append(
                    {
                        "cod_ibge": key[0],
                        "ano": key[1],
                        "field": field,
                        "type": "candidate_missing_value",
                        "reference": ref,
                        "candidate": None,
                    }
                )
                continue

            tolerance = (
                percentage_tolerance if field in PERCENT_FIELDS else money_tolerance
            )
            delta = abs(float(cand) - float(ref))
            if delta <= tolerance:
                exact_or_tolerant += 1
            else:
                mismatches.append(
                    {
                        "cod_ibge": key[0],
                        "ano": key[1],
                        "field": field,
                        "type": "value_difference",
                        "reference": ref,
                        "candidate": cand,
                        "delta": delta,
                        "tolerance": tolerance,
                    }
                )

    result = {
        "reference_municipalities": len(codes),
        "years": years,
        "reference_keys": len(reference),
        "comparisons": comparisons,
        "matched_within_tolerance": exact_or_tolerant,
        "both_absent": both_absent,
        "mismatch_count": len(mismatches),
        "strict_pass": len(mismatches) == 0,
        "mismatches": mismatches,
    }
    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Validate Gate E RGF history parity.")
    parser.add_argument("--candidate", type=Path, required=True)
    parser.add_argument("--reference-a01", type=Path, required=True)
    parser.add_argument("--reference-a02", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    result = compare_history(
        args.candidate,
        args.reference_a01,
        args.reference_a02,
        args.output,
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["strict_pass"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
