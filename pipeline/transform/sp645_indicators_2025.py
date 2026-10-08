from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INDICATORS = [
    "receita_tributaria_pct_receita_corrente",
    "iptu_pc",
    "itbi_pc",
    "iss_pc",
    "fpm_pct_receita_corrente",
    "icms_pct_receita_corrente",
    "ipva_pct_receita_corrente",
    "tributos_imobiliarios_pct_receita_tributaria",
    "transferencias_selecionadas_pct_receita_corrente",
    "investimento_pct_receita_corrente",
    "investimento_pc",
    "despesa_territorial_pct_despesa",
    "despesa_territorial_pc",
    "urbanismo_pct_territorial",
    "habitacao_pct_territorial",
    "saneamento_pct_territorial",
    "gestao_ambiental_pct_territorial",
    "transporte_pct_territorial",
    "dtp_pct_rcl",
    "dc_pct_rcl",
    "dcl_pct_rcl",
    "caixa_pos_rpnp_pct_rcl",
    "indicador_1",
    "indicador_2",
    "indicador_3",
]

TERRITORIAL_FIELDS = [
    "dca_func_urbanismo_liquidada",
    "dca_func_habitacao_liquidada",
    "dca_func_saneamento_liquidada",
    "dca_func_gestao_ambiental_liquidada",
    "dca_func_transporte_liquidada",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _number(value):
    if value in (None, ""):
        return None
    return float(value)


def _ratio(numerator, denominator, multiplier: float = 100.0):
    if numerator is None or denominator is None or denominator == 0:
        return None
    return numerator / denominator * multiplier


def _strict_sum(values):
    if any(value is None for value in values):
        return None
    return sum(values)


def calculate_row(row: dict) -> dict:
    rc = _number(row.get("dca_receita_corrente_bruta"))
    rt = _number(row.get("dca_receita_tributaria_bruta"))
    pop = _number(row.get("populacao_dca"))
    iptu = _number(row.get("dca_iptu_principal"))
    itbi = _number(row.get("dca_itbi_principal"))
    iss = _number(row.get("dca_iss_principal"))
    fpm = _number(row.get("dca_fpm_cota_mensal"))
    icms = _number(row.get("dca_icms_cota_parte"))
    ipva = _number(row.get("dca_ipva_cota_parte"))
    investimento = _number(row.get("dca_investimentos_liquidada"))
    despesa_total = _number(row.get("dca_despesa_total_liquidada"))

    territorial_values = [_number(row.get(field)) for field in TERRITORIAL_FIELDS]
    territorial = _strict_sum(territorial_values)

    result = {
        "receita_tributaria_pct_receita_corrente": _ratio(rt, rc),
        "iptu_pc": _ratio(iptu, pop, 1.0),
        "itbi_pc": _ratio(itbi, pop, 1.0),
        "iss_pc": _ratio(iss, pop, 1.0),
        "fpm_pct_receita_corrente": _ratio(fpm, rc),
        "icms_pct_receita_corrente": _ratio(icms, rc),
        "ipva_pct_receita_corrente": _ratio(ipva, rc),
        "tributos_imobiliarios_pct_receita_tributaria": _ratio(
            None if iptu is None or itbi is None else iptu + itbi,
            rt,
        ),
        "transferencias_selecionadas_pct_receita_corrente": _ratio(
            None if any(v is None for v in (fpm, icms, ipva)) else fpm + icms + ipva,
            rc,
        ),
        "investimento_pct_receita_corrente": _ratio(investimento, rc),
        "investimento_pc": _ratio(investimento, pop, 1.0),
        "despesa_territorial_pct_despesa": _ratio(territorial, despesa_total),
        "despesa_territorial_pc": _ratio(territorial, pop, 1.0),
        "dtp_pct_rcl": _number(row.get("rgf_dtp_percentual_rcl")),
        "dc_pct_rcl": _number(row.get("rgf02_dc_percentual_rcl")),
        "dcl_pct_rcl": _number(row.get("rgf02_dcl_percentual_rcl")),
        "caixa_pos_rpnp_pct_rcl": _ratio(
            _number(row.get("rgf_caixa_liquida_apos_rpnp")),
            _number(row.get("rreo_rcl_oficial")),
        ),
        "indicador_1": _number(row.get("indicador_1")),
        "indicador_2": _number(row.get("indicador_2")),
        "indicador_3": _number(row.get("indicador_3")),
    }

    labels = [
        ("urbanismo", "dca_func_urbanismo_liquidada"),
        ("habitacao", "dca_func_habitacao_liquidada"),
        ("saneamento", "dca_func_saneamento_liquidada"),
        ("gestao_ambiental", "dca_func_gestao_ambiental_liquidada"),
        ("transporte", "dca_func_transporte_liquidada"),
    ]
    for label, field in labels:
        result[f"{label}_pct_territorial"] = _ratio(
            _number(row.get(field)), territorial
        )

    return result


def calculate_workbook(source: Path, output: Path, manifest_path: Path) -> dict:
    source_book = openpyxl.load_workbook(source, read_only=True, data_only=True)
    source_sheet = source_book["Base multifuentes"]
    header = [cell.value for cell in next(source_sheet.iter_rows(min_row=1, max_row=1))]
    rows = []
    for values in source_sheet.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(header, values)))

    if len(rows) != 645:
        raise ValueError(f"Expected 645 source rows, got {len(rows)}")
    if len({str(row["cod_ibge"]) for row in rows}) != 645:
        raise ValueError("Municipality codes are not unique.")

    calculated = []
    for row in rows:
        calculated.append(
            {
                "cod_ibge": str(row["cod_ibge"]),
                "municipio": row.get("municipio"),
                "ano": int(row["ano"]),
                **calculate_row(row),
            }
        )

    coverage = {
        indicator: sum(row[indicator] is not None for row in calculated)
        for indicator in INDICATORS
    }

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Indicadores"
    output_fields = ["cod_ibge", "municipio", "ano", *INDICATORS]
    sheet.append(output_fields)
    for row in calculated:
        sheet.append([row.get(field) for field in output_fields])

    coverage_sheet = workbook.create_sheet("Cobertura")
    coverage_sheet.append(["indicador", "calculados", "ausentes", "cobertura_pct"])
    for indicator in INDICATORS:
        count = coverage[indicator]
        coverage_sheet.append([indicator, count, 645 - count, count / 645 * 100])

    method_sheet = workbook.create_sheet("Metodologia")
    method_sheet.append(["regra", "valor"])
    rules = [
        ("universo", "SP_645"),
        ("ano", "2025"),
        ("ausencia", "Nunca convertida em zero"),
        ("territorial", "Soma estrita das cinco funções; exige todos os componentes"),
        ("dtp_pct_rcl", "Percentual oficial RGF Anexo 01; não recomputado"),
        ("dc_pct_rcl", "Percentual oficial RGF Anexo 02; não recomputado"),
        ("dcl_pct_rcl", "Percentual oficial RGF Anexo 02; não recomputado"),
        ("capag", "Indicadores 1–3 oficiais; não recomputados"),
    ]
    for rule in rules:
        method_sheet.append(rule)

    fill = PatternFill("solid", fgColor="1F4E78")
    for current_sheet in (sheet, coverage_sheet, method_sheet):
        for cell in current_sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        current_sheet.freeze_panes = "A2"
        current_sheet.auto_filter.ref = current_sheet.dimensions

    for index, field in enumerate(output_fields, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(
            max(12, len(field) + 2), 34
        )
    sheet.column_dimensions["B"].width = 28
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "@"
        for cell in row[3:]:
            cell.number_format = "0.00"
    for row in coverage_sheet.iter_rows(min_row=2):
        row[3].number_format = "0.00"

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)

    result = {
        "rows": 645,
        "indicators": len(INDICATORS),
        "input_sha256": sha256(source),
        "output_sha256": sha256(output),
        "coverage": coverage,
        "coverage_pct": {
            key: round(value / 645 * 100, 2) for key, value in coverage.items()
        },
    }
    manifest_path.parent.mkdir(parents=True, exist_ok=True)
    manifest_path.write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description="Calculate SP645 2025 Gate D indicators.")
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    args = parser.parse_args()
    print(
        json.dumps(
            calculate_workbook(args.source, args.output, args.manifest),
            ensure_ascii=False,
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
