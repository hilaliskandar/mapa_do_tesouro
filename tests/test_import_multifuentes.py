from pathlib import Path
import sqlite3

import openpyxl
import yaml

from pipeline.ingest.import_multifuentes import import_multifuentes

ROOT = Path(__file__).resolve().parents[1]


def create_fixture(path: Path, mapping_path: Path) -> None:
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    headers = list(mapping["keys"].values())
    for spec in mapping["variables"].values():
        if spec["column"] not in headers:
            headers.append(spec["column"])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = mapping["source"]["sheet"]
    ws.append(headers)

    for year in range(2013, 2026):
        row = {header: None for header in headers}
        row["cod_ibge"] = "3500000"
        row["municipio"] = "Municipio Teste"
        row["ano"] = year
        for variable_id, spec in mapping["variables"].items():
            if spec["availability_start"] <= year <= spec["availability_end"]:
                if spec["value_type"] == "numeric":
                    row[spec["column"]] = float(year)
                else:
                    row[spec["column"]] = "A"
        ws.append([row[h] for h in headers])

    wb.save(path)


def test_importer_preserves_applicability_and_observed_zero(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )

    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    database = tmp_path / "fixture.sqlite"
    result = import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        overwrite=False,
        build_timestamp="2026-10-06T00:00:00-03:00",
    )

    assert result["rows"] == 13
    assert result["municipalities"] == 1

    con = sqlite3.connect(database)
    try:
        assert con.execute(
            """
            SELECT status FROM observacao
            WHERE codigo_ibge='3500000'
              AND ano=2013
              AND variavel_id='rreo_rcl_oficial'
            """
        ).fetchone()[0] == "nao_aplicavel"

        assert con.execute(
            """
            SELECT status FROM observacao
            WHERE codigo_ibge='3500000'
              AND ano=2015
              AND variavel_id='rreo_rcl_oficial'
            """
        ).fetchone()[0] == "observado"

        assert con.execute(
            """
            SELECT status FROM observacao
            WHERE codigo_ibge='3500000'
              AND ano=2024
              AND variavel_id='capag'
            """
        ).fetchone()[0] == "nao_aplicavel"

        assert con.execute(
            """
            SELECT status, valor_texto FROM observacao
            WHERE codigo_ibge='3500000'
              AND ano=2025
              AND variavel_id='capag'
            """
        ).fetchone() == ("observado", "A")
    finally:
        con.close()


def test_percentage_scale_is_applied(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    wb = openpyxl.load_workbook(workbook)
    ws = wb[mapping["source"]["sheet"]]
    headers = [c.value for c in ws[1]]
    col = headers.index(mapping["variables"]["dtp_pct_rcl"]["column"]) + 1
    ws.cell(row=4, column=col, value=36.69)
    wb.save(workbook)

    database = tmp_path / "fixture.sqlite"
    import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        build_timestamp="2026-10-06T00:00:00-03:00",
    )

    con = sqlite3.connect(database)
    try:
        value = con.execute(
            """
            SELECT valor_num FROM observacao
            WHERE codigo_ibge='3500000'
              AND ano=2015
              AND variavel_id='dtp_pct_rcl'
            """
        ).fetchone()[0]
        assert abs(value - 0.3669) < 1e-12
    finally:
        con.close()


def test_importer_records_source_artifact_and_field_provenance(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    database = tmp_path / "fixture.sqlite"
    result = import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        build_timestamp="2026-10-06T00:00:00-03:00",
    )

    con = sqlite3.connect(database)
    try:
        artifact = con.execute(
            """
            SELECT fonte_id,nome,sha256,mime_type
            FROM artefato_fonte
            WHERE artefato_id=?
            """,
            (result["source_artifact_id"],),
        ).fetchone()
        assert artifact is not None
        assert artifact[0] == "PROJETO_BASE_MULTIFONTES"
        assert artifact[1] == workbook.name
        assert artifact[2] == result["source_sha256"]
        assert artifact[3] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

        provenance = con.execute(
            """
            SELECT tipo,artefato_id,origem_aba,origem_campo,origem_referencia
            FROM observacao_proveniencia
            WHERE codigo_ibge='3500000'
              AND ano=2015
              AND variavel_id='dtp_pct_rcl'
            """
        ).fetchone()
        assert provenance[0] == "campo_fonte"
        assert provenance[1] == result["source_artifact_id"]
        assert provenance[2] == mapping["source"]["sheet"]
        assert provenance[3] == mapping["variables"]["dtp_pct_rcl"]["column"]
        assert "row=4" in provenance[4]

        assert result["schema_version"] == "0.3.0"
        assert result["provenance_rows"] == (
            result["observations"]
        )
        assert con.execute(
            "SELECT schema_version FROM build WHERE build_id=?",
            (result["build_id"],),
        ).fetchone()[0] == "0.3.0"
    finally:
        con.close()


def test_importer_supports_custom_universe_from_mapping(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1
    mapping["universe"] = {
        "id": "SP_TESTE",
        "name": "São Paulo teste",
        "description": "Universo de teste.",
        "type": "analitico",
    }

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    database = tmp_path / "fixture.sqlite"
    result = import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        build_timestamp="2026-10-07T00:00:00-03:00",
    )

    assert result["universe_id"] == "SP_TESTE"
    assert result["universe_name"] == "São Paulo teste"

    con = sqlite3.connect(database)
    try:
        assert con.execute(
            "SELECT universo_id,nome FROM universo"
        ).fetchall() == [("SP_TESTE", "São Paulo teste")]
        assert con.execute(
            "SELECT DISTINCT universo_id FROM universo_municipio"
        ).fetchall() == [("SP_TESTE",)]
        assert con.execute(
            "SELECT DISTINCT universo_id FROM cobertura"
        ).fetchall() == [("SP_TESTE",)]
    finally:
        con.close()


def test_importer_treats_trailing_empty_cells_as_absent(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    wb = openpyxl.load_workbook(workbook)
    ws = wb[mapping["source"]["sheet"]]
    trailing_column = mapping["variables"]["indicador_3"]["column"]
    trailing_index = {
        cell.value: cell.column
        for cell in ws[1]
        if cell.value is not None
    }[trailing_column]
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=trailing_index).value = None
    wb.save(workbook)

    database = tmp_path / "trailing.sqlite"
    import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        build_timestamp="2026-10-07T00:00:00-03:00",
    )

    con = sqlite3.connect(database)
    try:
        assert con.execute(
            """
            SELECT COUNT(*)
            FROM observacao
            WHERE variavel_id='indicador_3'
              AND status IN ('ausente','nao_aplicavel')
            """
        ).fetchone()[0] == 13
    finally:
        con.close()



def test_importer_preserves_explicit_numeric_missing_token_as_absent(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1
    mapping["variables"]["indicador_3"]["missing_tokens"] = ["n.d."]

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    wb = openpyxl.load_workbook(workbook)
    ws = wb[mapping["source"]["sheet"]]
    headers = [cell.value for cell in ws[1]]
    col = headers.index(mapping["variables"]["indicador_3"]["column"]) + 1
    ws.cell(row=ws.max_row, column=col, value="n.d.")
    wb.save(workbook)

    database = tmp_path / "missing-token.sqlite"
    import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        build_timestamp="2026-10-09T00:00:00-03:00",
    )

    con = sqlite3.connect(database)
    try:
        status, value = con.execute(
            """
            SELECT status,valor_num
            FROM observacao
            WHERE codigo_ibge='3500000'
              AND ano=2025
              AND variavel_id='indicador_3'
            """
        ).fetchone()
        assert status == "ausente"
        assert value is None
    finally:
        con.close()



def test_importer_coverage_has_review_status_column(tmp_path):
    mapping_path = ROOT / "data" / "mappings" / "base_multifuentes_v0_4.yml"
    mapping = yaml.safe_load(mapping_path.read_text(encoding="utf-8"))
    mapping["source"]["expected_rows"] = 13
    mapping["source"]["expected_municipalities"] = 1

    local_mapping = tmp_path / "mapping.yml"
    local_mapping.write_text(
        yaml.safe_dump(mapping, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    workbook = tmp_path / "fixture.xlsx"
    create_fixture(workbook, local_mapping)

    database = tmp_path / "coverage.sqlite"
    import_multifuentes(
        workbook,
        database,
        mapping_path=local_mapping,
        build_timestamp="2026-10-09T00:00:00-03:00",
    )

    con = sqlite3.connect(database)
    try:
        row = con.execute(
            """
            SELECT observado,ausente,nao_aplicavel,em_revisao
            FROM cobertura
            WHERE variavel_id='capag' AND ano=2025
            """
        ).fetchone()
        assert row is not None
        assert row[3] == 0
    finally:
        con.close()
