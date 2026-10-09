from pathlib import Path

import openpyxl
import yaml

from pipeline.build.sp645_private_input import prepare_sp645_private_input


def _write_history(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Tipologia SP645"
    fields = [
        "cod_ibge",
        "municipio",
        "ano",
        "receita_tributaria_pct_receita_corrente",
        "investimento_pct_receita_corrente",
        "despesa_territorial_pct_despesa",
        "dtp_pct_rcl",
        "dc_pct_rcl",
        "dcl_pct_rcl",
    ]
    sheet.append(fields)
    for index in range(645):
        code = f"{3500000 + index:07d}"
        for year in range(2021, 2026):
            sheet.append([code, f"M{index}", year, 10, 5, 20, 40 if year >= 2023 else None, 8 if year >= 2023 else None, 2 if year >= 2023 else None])
    book.save(path)


def _write_current(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Base multifuentes"
    sheet.append(["cod_ibge", "municipio", "ano", "dca_receita_corrente_bruta"])
    for index in range(645):
        code = f"{3500000 + index:07d}"
        sheet.append([code, f"M{index}", 2025, 1000 + index])
    book.save(path)


def test_prepare_sp645_private_input_preserves_history_and_scopes_raw_2025(tmp_path):
    history = tmp_path / "history.xlsx"
    current = tmp_path / "current.xlsx"
    output = tmp_path / "combined.xlsx"
    mapping = tmp_path / "base.yml"
    mapping_out = tmp_path / "mapping.yml"
    manifest = tmp_path / "manifest.json"

    _write_history(history)
    _write_current(current)
    mapping.write_text(
        yaml.safe_dump(
            {
                "universe": {"id": "SP_645"},
                "source": {
                    "source_id": "PROJETO_BASE_MULTIFONTES_SP645_2025",
                    "workbook_name": "current.xlsx",
                    "sheet": "Base multifuentes",
                    "data_version": "x",
                    "expected_rows": 645,
                    "expected_municipalities": 645,
                    "expected_years": {"start": 2025, "end": 2025},
                },
                "keys": {"codigo_ibge": "cod_ibge", "municipio": "municipio", "ano": "ano"},
                "variables": {
                    "dtp_pct_rcl": {"column": "rgf_dtp_percentual_rcl", "source_id": "SICONFI_RGF_A01", "availability_start": 2025, "availability_end": 2025, "value_type": "numeric", "scale": 0.01},
                    "dc_pct_rcl": {"column": "rgf02_dc_percentual_rcl", "source_id": "SICONFI_RGF_A02", "availability_start": 2025, "availability_end": 2025, "value_type": "numeric", "scale": 0.01},
                    "dcl_pct_rcl": {"column": "rgf02_dcl_percentual_rcl", "source_id": "SICONFI_RGF_A02", "availability_start": 2025, "availability_end": 2025, "value_type": "numeric", "scale": 0.01},
                },
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )

    result = prepare_sp645_private_input(
        history, current, mapping, output, mapping_out, manifest
    )
    assert result["rows"] == 3225
    assert result["municipalities"] == 645
    assert result["years"] == [2021, 2022, 2023, 2024, 2025]

    book = openpyxl.load_workbook(output, read_only=True, data_only=True)
    try:
        sheet = book["SP645 Site Input"]
        rows = list(sheet.iter_rows(values_only=True))
        header = list(rows[0])
        idx = {name: i for i, name in enumerate(header)}
        row_2024 = rows[1 + 3]
        row_2025 = rows[1 + 4]
        assert row_2024[idx["ano"]] == 2024
        assert row_2024[idx["dca_receita_corrente_bruta"]] in (None, "")
        assert row_2025[idx["dca_receita_corrente_bruta"]] == 1000
    finally:
        book.close()

    generated = yaml.safe_load(mapping_out.read_text(encoding="utf-8"))
    assert generated["source"]["expected_rows"] == 3225
    assert generated["source"]["expected_years"] == {"start": 2021, "end": 2025}
    assert generated["variables"]["dtp_pct_rcl"]["availability_start"] == 2023
    assert generated["variables"]["receita_tributaria_pct_receita_corrente"]["scale"] == 0.01
